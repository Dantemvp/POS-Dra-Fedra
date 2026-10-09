"""Cotejo local de hojas LookinBody contra Excel y pacientes. No escribe en la BD.

Genera PDFs contenedores de los JPEG originales, sin OCR/IA ni reinterpretación.
Los datos y resultados deben permanecer fuera de cualquier repositorio Git.
"""
import argparse
from collections import Counter, defaultdict
from datetime import datetime
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import subprocess

from PIL import Image
import openpyxl
from reportlab.pdfgen import canvas

spec = importlib.util.spec_from_file_location("historicos", Path(__file__).with_name("importar-documentos-historicos.py"))
historicos = importlib.util.module_from_spec(spec)
spec.loader.exec_module(historicos)

def huella(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()

def clave_archivo(nombre):
    m = re.fullmatch(r"(.+?)_(\d{14})_InBody\.jpg", nombre, re.I)
    if not m:
        raise ValueError("Nombre de hoja no reconocido.")
    return m[1], datetime.strptime(m[2], "%Y%m%d%H%M%S")

def cotejar(archivos, filas, pacientes):
    estudios = defaultdict(list)
    nombres = defaultdict(list)
    for paciente in pacientes:
        nombres[historicos.normalizar(paciente["nombre"] + " " + (paciente.get("apellidos") or ""))].append(paciente)
    for fila in filas:
        if not any(v is not None for v in fila):
            continue
        instante = datetime.strptime(str(fila[13]), "%Y.%m.%d. %H:%M:%S")
        estudios[(str(fila[1]), instante)].append(fila)
    resultados = []
    cubiertos = set()
    for archivo in archivos:
        r = {"archivo": str(archivo), "estado": "pendiente"}
        try:
            clave = clave_archivo(archivo.name)
        except ValueError:
            r["estado"] = "nombre_archivo_no_reconocido"
            resultados.append(r)
            continue
        rows = estudios.get(clave, [])
        cubiertos.add(clave)
        r.update(id_lookinbody=clave[0], fecha_hora_local=clave[1].isoformat())
        if len(rows) != 1:
            r["estado"] = "estudio_excel_ambiguo" if rows else "sin_estudio_excel"
        else:
            nombre = str(rows[0][0]).strip()
            ps = nombres[historicos.normalizar(nombre)]
            r["nombre_excel"] = nombre
            if len(ps) == 1:
                r.update(estado="listo", paciente_id=ps[0]["id"])
            else:
                r["estado"] = "nombre_ambiguo" if ps else "sin_paciente"
        resultados.append(r)
    faltantes = [{"id_lookinbody":k[0], "fecha_hora_local":k[1].isoformat()}
        for k in sorted(set(estudios) - cubiertos)]
    return resultados, faltantes

def pdf_de_jpg(imagen, salida):
    with Image.open(imagen) as img:
        if img.format != "JPEG":
            raise ValueError("La hoja no es un JPEG válido.")
        ancho, alto = img.size
        img.verify()
    # Página con la proporción exacta de la hoja. drawImage incrusta el JPEG
    # sin volver a comprimirlo; invariant evita cambios de hash entre corridas.
    ancho_pt = 595.276
    alto_pt = ancho_pt * alto / ancho
    c = canvas.Canvas(str(salida), pagesize=(ancho_pt, alto_pt), invariant=1)
    c.setTitle("Hoja InBody histórica")
    c.drawImage(str(imagen), 0, 0, width=ancho_pt, height=alto_pt)
    c.showPage()
    c.save()

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--images", required=True, type=Path)
    p.add_argument("--excel", required=True, type=Path)
    p.add_argument("--patients", required=True, type=Path)
    p.add_argument("--output", required=True, type=Path)
    args = p.parse_args()
    output = args.output.resolve()
    probe = subprocess.run(["git", "-C", str(output.parent), "rev-parse", "--show-toplevel"], capture_output=True, text=True)
    if probe.returncode == 0 or output.is_relative_to(Path(__file__).resolve().parents[2]):
        p.error("La salida clínica debe permanecer fuera de Git.")
    if output.exists():
        p.error("La salida ya existe. No sobrescribir resultados.")
    imagenes = sorted(args.images.resolve().rglob("*.jpg"))
    if not imagenes:
        p.error("No hay hojas JPG.")
    hash_excel = huella(args.excel)
    hash_padron = huella(args.patients)
    pacientes = json.loads(args.patients.read_text(encoding="utf-8-sig"))
    wb = openpyxl.load_workbook(args.excel, read_only=True, data_only=True)
    try:
        sheet = wb["InBody"]
        headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        if headers[0:2] != ("1. Name", "2. ID") or headers[13] != "14. Test Date / Time":
            raise ValueError("El Excel no tiene las columnas LookinBody esperadas.")
        resultados, faltantes = cotejar(imagenes, sheet.iter_rows(min_row=2, values_only=True), pacientes)
    finally:
        wb.close()
    if huella(args.excel) != hash_excel or huella(args.patients) != hash_padron:
        raise ValueError("Una fuente cambió durante el cotejo.")
    output.mkdir(parents=True)
    pdfs = output / "pdf"
    pdfs.mkdir()
    registros = []
    vistos = set()
    for r in resultados:
        jpg = Path(r["archivo"])
        r["sha256_jpg"] = huella(jpg)
        if r["estado"] != "listo":
            continue
        clave = (r["paciente_id"], r["sha256_jpg"])
        if clave in vistos:
            r["estado"] = "duplicado"
            continue
        vistos.add(clave)
        carpeta = pdfs / r["paciente_id"]
        carpeta.mkdir(exist_ok=True)
        pdf = carpeta / (r["fecha_hora_local"].replace(":", "-") + "_" + r["sha256_jpg"] + ".pdf")
        pdf_de_jpg(jpg, pdf)
        if huella(jpg) != r["sha256_jpg"]:
            raise ValueError("La imagen cambió durante la conversión.")
        registros.append({"estado":"listo_documento", "paciente_id":r["paciente_id"],
            "nombre_interno":r["nombre_excel"], "archivo_relativo":str(pdf.relative_to(pdfs)),
            "sha256":huella(pdf), "sha256_jpg":r["sha256_jpg"], "id_lookinbody":r["id_lookinbody"],
            "fecha":r["fecha_hora_local"][:10], "fecha_hora_local":r["fecha_hora_local"]})
    manifiesto = {"criterio_vinculacion":"nombre_unico", "tipo_documento":"inbody",
        "origen_datos":"lookinbody_imagenes_20261008", "origen":str(pdfs),
        "pacientes_fuente":str(args.patients.resolve()), "pacientes_sha256":hash_padron,
        "excel_fuente":str(args.excel.resolve()), "excel_sha256":hash_excel, "recetas":registros}
    resumen = {"imagenes":len(imagenes), "estados":dict(Counter(r["estado"] for r in resultados)),
        "pdfs_vinculables":len(registros), "pacientes_vinculables":len(set(r["paciente_id"] for r in registros)),
        "estudios_sin_imagen":len(faltantes)}
    (output/"manifest.json").write_text(json.dumps(manifiesto, ensure_ascii=False, indent=2), encoding="utf-8")
    (output/"cotejo.json").write_text(json.dumps({"resumen":resumen,"hojas":resultados,"faltantes":faltantes}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(resumen))

if __name__ == "__main__":
    main()
