"""Normaliza un conteo físico de inventario sin tocar la base de datos.

Este script es deliberadamente de solo lectura. Produce un resumen en pantalla
y, si se solicita, un JSON intermedio que puede revisarse antes de preparar una
conciliación con producción.

Uso:
    python scripts/inventario-conteo-fisico.py archivo.xlsx
    python scripts/inventario-conteo-fisico.py archivo.xlsx --salida conteo.json
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import unicodedata
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl


HOJAS = {
    "INVENTARIO (BODEGA)": "Bodega",
    "INVENTARIO (FARMACIA)": "Farmacia",
}

ENCABEZADOS = [
    "PRODUCTO",
    "CONCEPTO",
    "PRESENTACION",
    "LOTE",
    "FX. VENCIMIENTO",
    "STATUS",
    "CANTIDAD INV. (PIEZAS)",
    "PRECIO UNITARIO (Venta)",
    "COSTO TOTAL INV.",
    "UBICACION",
]


def texto(valor: Any) -> str | None:
    if valor is None:
        return None
    limpio = re.sub(r"\s+", " ", str(valor)).strip()
    if not limpio or limpio.casefold() in {"n/a", "na"}:
        return None
    return limpio


def clave_texto(valor: str | None) -> str:
    base = unicodedata.normalize("NFKD", valor or "")
    sin_acentos = "".join(c for c in base if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", sin_acentos.casefold()).strip()


def numero(valor: Any) -> Decimal | None:
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float, Decimal)):
        if isinstance(valor, float) and not math.isfinite(valor):
            return None
        return Decimal(str(valor))
    limpio = texto(valor)
    if limpio is None:
        return None
    try:
        return Decimal(limpio.replace(",", ""))
    except Exception:
        return None


def fecha_iso(valor: Any) -> str | None:
    if isinstance(valor, (datetime, date)):
        return valor.date().isoformat() if isinstance(valor, datetime) else valor.isoformat()
    limpio = texto(valor)
    return limpio


def serializa(valor: Any) -> Any:
    if isinstance(valor, Decimal):
        return float(valor) if valor % 1 else int(valor)
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    return valor


def validar_encabezados(nombre_hoja: str, hoja: Any) -> None:
    recibidos = [celda.value for celda in hoja[3][: len(ENCABEZADOS)]]
    esperados = [clave_texto(valor) for valor in ENCABEZADOS]
    actuales = [clave_texto(texto(valor)) for valor in recibidos]
    diferencias = [
        {
            "columna": indice + 1,
            "esperado": ENCABEZADOS[indice],
            "recibido": recibidos[indice],
        }
        for indice, (esperado, actual) in enumerate(zip(esperados, actuales))
        if esperado != actual
    ]
    if diferencias:
        raise ValueError(
            f"La hoja {nombre_hoja!r} no tiene el formato esperado: "
            + json.dumps(diferencias, ensure_ascii=False, default=serializa)
        )


def misma_ruta(origen: Path, salida: Path) -> bool:
    if origen.resolve() == salida.resolve():
        return True
    return salida.exists() and os.path.samefile(origen, salida)


def leer(archivo: Path) -> dict[str, Any]:
    formulas = openpyxl.load_workbook(archivo, read_only=True, data_only=False)
    valores = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    faltantes = [nombre for nombre in HOJAS if nombre not in valores.sheetnames]
    if faltantes:
        raise ValueError(f"Faltan hojas requeridas: {', '.join(faltantes)}")

    filas: list[dict[str, Any]] = []
    bloqueos: list[dict[str, Any]] = []
    avisos: list[dict[str, Any]] = []

    for nombre_hoja, ubicacion_esperada in HOJAS.items():
        hoja_f = formulas[nombre_hoja]
        hoja_v = valores[nombre_hoja]
        validar_encabezados(nombre_hoja, hoja_f)
        for renglon, (fila_f, fila_v) in enumerate(
            zip(
                hoja_f.iter_rows(min_row=4, values_only=True),
                hoja_v.iter_rows(min_row=4, values_only=True),
            ),
            start=4,
        ):
            if not any(v is not None for v in fila_v[:10]):
                continue

            marca = texto(fila_v[0])
            concepto = texto(fila_v[1])
            presentacion = texto(fila_v[2])
            lote = texto(fila_v[3])
            caducidad = fecha_iso(fila_v[4])
            cantidad = numero(fila_v[6])
            precio = numero(fila_v[7])
            ubicacion = texto(fila_v[9]) or ubicacion_esperada
            nombre_producto = " - ".join(
                parte for parte in (marca, concepto, presentacion) if parte
            )
            referencia = {"hoja": nombre_hoja, "renglon": renglon}

            if not nombre_producto:
                bloqueos.append({**referencia, "motivo": "producto sin nombre"})
                continue
            if cantidad is None:
                bloqueos.append({
                    **referencia,
                    "producto": nombre_producto,
                    "motivo": "cantidad vacía o no numérica",
                    "valor": fila_f[6],
                })
            elif cantidad < 0:
                bloqueos.append({
                    **referencia,
                    "producto": nombre_producto,
                    "motivo": "cantidad negativa",
                    "valor": serializa(cantidad),
                })

            if cantidad and not lote:
                avisos.append({**referencia, "producto": nombre_producto, "motivo": "existencia positiva sin lote"})
            if cantidad and not caducidad:
                avisos.append({**referencia, "producto": nombre_producto, "motivo": "existencia positiva sin caducidad"})
            if precio is None or precio == 0:
                avisos.append({**referencia, "producto": nombre_producto, "motivo": "precio vacío o en cero"})

            filas.append(
                {
                    "origen_hoja": nombre_hoja,
                    "origen_renglon": renglon,
                    "marca": marca,
                    "concepto": concepto,
                    "presentacion": presentacion,
                    "nombre_producto": nombre_producto,
                    "clave_producto": clave_texto(nombre_producto),
                    "lote": lote,
                    "caducidad": caducidad,
                    "cantidad": serializa(cantidad),
                    "precio_venta": serializa(precio),
                    "ubicacion": ubicacion,
                }
            )

    claves_lote = Counter(
        (fila["clave_producto"], clave_texto(fila["lote"]), clave_texto(fila["ubicacion"]))
        for fila in filas
    )
    duplicados = [
        {"clave_producto": clave[0], "lote": clave[1] or None, "ubicacion": clave[2], "repeticiones": total}
        for clave, total in claves_lote.items()
        if total > 1
    ]
    productos = {fila["clave_producto"] for fila in filas}
    piezas = sum(Decimal(str(fila["cantidad"] or 0)) for fila in filas if fila["cantidad"] is not None)
    valor = sum(
        Decimal(str(fila["cantidad"] or 0)) * Decimal(str(fila["precio_venta"] or 0))
        for fila in filas
        if fila["cantidad"] is not None and fila["precio_venta"] is not None
    )

    return {
        "archivo": str(archivo.resolve()),
        "sha256": hashlib.sha256(archivo.read_bytes()).hexdigest(),
        "modo": "solo lectura; no concilia ni escribe en producción",
        "resumen": {
            "renglones_lote": len(filas),
            "productos_normalizados": len(productos),
            "piezas": serializa(piezas),
            "valor_venta_estimado": serializa(valor),
            "bloqueos": len(bloqueos),
            "avisos": len(avisos),
            "claves_producto_lote_ubicacion_duplicadas": len(duplicados),
        },
        "bloqueos": bloqueos,
        "avisos": avisos,
        "duplicados": duplicados,
        "filas": filas,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("archivo", type=Path)
    parser.add_argument("--salida", type=Path)
    args = parser.parse_args()
    if not args.archivo.is_file():
        parser.error(f"No existe el archivo: {args.archivo}")
    if args.salida and misma_ruta(args.archivo, args.salida):
        parser.error("La salida no puede ser el mismo archivo que la fuente.")

    try:
        resultado = leer(args.archivo)
    except ValueError as error:
        parser.error(str(error))
    print(json.dumps(resultado["resumen"], ensure_ascii=False, indent=2))
    if resultado["bloqueos"]:
        print("\nBloqueos:")
        print(json.dumps(resultado["bloqueos"], ensure_ascii=False, indent=2, default=serializa))
    if args.salida:
        args.salida.parent.mkdir(parents=True, exist_ok=True)
        args.salida.write_text(
            json.dumps(resultado, ensure_ascii=False, indent=2, default=serializa),
            encoding="utf-8",
        )
        print(f"\nSalida de revisión: {args.salida.resolve()}")
    return 2 if resultado["bloqueos"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
